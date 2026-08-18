"""Tests for the Network Infrastructure SOD Report.

Its own module rather than another thousand lines onto tests.py: this report has a distinct
failure mode from the other two — it is mostly BLANKS by design, and the thing worth guarding
is that blank stays blank and never quietly becomes an all-clear.
"""
from unittest import mock

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.test import TestCase
from django.urls import reverse

from . import network_sod as ns


class SodChecklistShape(TestCase):
    """The catalogue must match the approved reference sheet — the two are compared side by
    side every morning, so a section that drifted would be found the hard way."""

    def test_counts_match_the_reference_sheet(self):
        summary = ns.summarise(ns.blank_checklist())
        self.assertEqual(summary["core_wan_count"], 6)
        self.assertEqual(summary["firewall_count"], 5)
        self.assertEqual(summary["circuit_count"], 4)
        self.assertEqual(summary["controller_count"], 2)
        self.assertEqual(summary["waf_count"], 18)
        self.assertEqual(summary["ping_total"], 11)

    def test_a_blank_checklist_raises_no_warnings(self):
        """An unfilled sheet must not manufacture an all-clear OR a fault."""
        summary = ns.summarise(ns.blank_checklist())
        self.assertEqual(summary["banners"], [])
        self.assertEqual(summary["links_down"], 0)
        self.assertEqual(summary["degraded_circuits"], 0)
        self.assertEqual(summary["rogue_controllers"], 0)

    def test_every_key_is_unique(self):
        """Keys name the form fields. A duplicate would silently overwrite an answer."""
        data = ns.blank_checklist()
        keys = ([c.key for c in data["core_wan"]]
                + [c.key for g in data["firewalls"] for c in g.checks]
                + [c.key for c in data["floor"]])
        self.assertEqual(len(keys), len(set(keys)))


class SodDegradationRules(TestCase):
    def test_an_unmeasured_circuit_is_not_reported_as_degraded(self):
        """An em dash or "Test incomplete" means the test produced no answer, which is a
        different claim from the circuit being bad. The reference sheet draws the line in
        exactly this place: it flags Telecontract, whose loss test failed, but not
        Wafanyakazi, whose test simply did not run."""
        data = ns.blank_checklist()
        by_key = {c.key: c for c in data["circuits"]}
        by_key["wafanyakazi"].loss = "—"
        by_key["wafanyakazi"].quality = "Test incomplete"
        by_key["telecontract"].loss = "Failed"
        by_key["dandemutande"].loss = "5.70%"
        by_key["dandemutande"].quality = "Average/Poor/Average"

        summary = ns.summarise(data)
        self.assertEqual(summary["degraded_circuits"], 2)
        # The banner names the LINK ("Dandemutande Internet"), matching how the WLAN banner
        # says "Harare WLAN Controller"; the table below names the provider.
        flagged = [r["name"] for r in summary["banners"][0]["rows"]]
        self.assertEqual(flagged, ["Dandemutande Internet", "Telecontract Internet"],
                         "worst first: a measured 5.70% + Poor outranks a failed test")

    def test_a_clean_zero_loss_is_not_degraded(self):
        data = ns.blank_checklist()
        data["circuits"][0].loss = "0%"
        data["circuits"][0].quality = "Good/Good/Good"
        self.assertEqual(ns.summarise(data)["degraded_circuits"], 0)

    def test_a_failed_loss_test_is_worded_as_a_failure_not_a_reading(self):
        data = ns.blank_checklist()
        data["circuits"][0].loss = "Failed"
        row = ns.summarise(data)["banners"][0]["rows"][0]
        self.assertIn("failed to complete", row["detail"])
        self.assertNotIn("Packet loss Failed", row["detail"])

    def test_a_failure_cause_reaches_the_banner_but_not_the_narrow_column(self):
        """"Failed (connection timed out)" carries its cause into the banner sentence, while
        the 7-wide Pkt Loss cell keeps just "Failed" — there is no room for more there, and
        the reference sheet splits it exactly this way."""
        data = ns.blank_checklist()
        data["circuits"][0].loss = "Failed (connection timed out)"
        detail = ns.summarise(data)["banners"][0]["rows"][0]["detail"]
        self.assertEqual(detail,
                         "Packet-loss measurement failed to complete (connection timed out)")
        self.assertEqual(ns.loss_short("Failed (connection timed out)"), "Failed")
        self.assertEqual(ns.loss_short("5.70%"), "5.70%")

    def test_at_most_two_quality_aspects_reach_the_banner(self):
        """The banner is a headline; the Quality column below carries the full triple."""
        data = ns.blank_checklist()
        data["circuits"][0].loss = "1%"
        data["circuits"][0].quality = "Average/Poor/Average"
        detail = ns.summarise(data)["banners"][0]["rows"][0]["detail"]
        self.assertEqual(detail.count("  ·  "), 2)      # loss + two aspects
        self.assertNotIn("Chat", detail)

    def test_quality_triple_is_decomposed_worst_first(self):
        data = ns.blank_checklist()
        data["circuits"][0].quality = "Average/Poor/Average"
        detail = ns.summarise(data)["banners"][0]["rows"][0]["detail"]
        self.assertIn("Online Gaming: Poor", detail)
        self.assertLess(detail.index("Online Gaming: Poor"),
                        detail.index("Video Streaming: Average"))

    def test_free_text_quality_is_passed_through_not_split(self):
        data = ns.blank_checklist()
        data["circuits"][0].loss = "1%"
        data["circuits"][0].quality = "Test incomplete"
        detail = ns.summarise(data)["banners"][0]["rows"][0]["detail"]
        self.assertIn("Test incomplete", detail)

    def test_rogue_aps_raise_a_warning_per_controller(self):
        data = ns.blank_checklist()
        data["controllers"][0].rogue = "592"
        summary = ns.summarise(data)
        self.assertEqual(summary["rogue_controllers"], 1)
        self.assertTrue(any("ROGUE ACCESS POINTS" in b["headline"]
                            for b in summary["banners"]))

    def test_a_non_numeric_rogue_count_does_not_explode(self):
        data = ns.blank_checklist()
        data["controllers"][0].rogue = "none seen"
        self.assertEqual(ns.summarise(data)["rogue_controllers"], 0)

    def test_a_down_check_counts_as_a_link_down(self):
        data = ns.blank_checklist()
        data["core_wan"][0].status = "DOWN"
        self.assertEqual(ns.summarise(data)["links_down"], 1)


class SodWorkbookRendersEverywhere(TestCase):
    def test_every_colour_is_opaque(self):
        """Colours must carry an FF alpha, not 00.

        The engine stores its palette as 00RRGGBB. Excel ignores that leading pair, but
        LibreOffice, Google Sheets and several web previewers read 00 as fully transparent
        and drop every fill and font colour — rendering this dark report as unstyled
        black-on-white that looks nothing like the approved sheet. This is the regression
        that produced exactly that.

        Only colours actually SET are checked. openpyxl reports an unstyled cell's fill as
        the default "00000000", and the merge tails legitimately carry that — the reference
        sheet leaves them unfilled too, because Excel paints a merged range from its anchor.
        """
        import io

        import openpyxl

        payload = ns.build_report(ns.blank_checklist(), theme="dark", author="A")
        ws = openpyxl.load_workbook(io.BytesIO(payload)).active
        bad = []
        for row in ws.iter_rows(min_row=1, max_row=90, max_col=22):
            for cell in row:
                if cell.fill is not None and cell.fill.patternType == "solid":
                    rgb = getattr(cell.fill.fgColor, "rgb", None)
                    if isinstance(rgb, str) and not rgb.upper().startswith("FF"):
                        bad.append("{} fill {}".format(cell.coordinate, rgb))
                if cell.value is not None and cell.font is not None:
                    rgb = getattr(cell.font.color, "rgb", None)
                    if isinstance(rgb, str) and not rgb.upper().startswith("FF"):
                        bad.append("{} font {}".format(cell.coordinate, rgb))
        self.assertEqual(bad[:10], [], "{} non-opaque colours".format(len(bad)))

    def test_the_canvas_is_painted_to_the_reference_floor(self):
        """A quiet morning must not produce a visibly shorter page than a busy one."""
        import io

        import openpyxl

        payload = ns.build_report(ns.blank_checklist(), theme="dark", author="A")
        ws = openpyxl.load_workbook(io.BytesIO(payload)).active
        self.assertGreaterEqual(ws.max_row, 140)


class SodCollectIsHonestAboutWhatItKnows(TestCase):
    def test_prometheus_being_down_leaves_the_sheet_blank_rather_than_failing(self):
        """The engineer hand-keyed this sheet before the screen existed; an unreachable
        Prometheus must not stop them doing it again."""
        with mock.patch("reports.network_sod.network.device_inventory",
                        side_effect=RuntimeError("prometheus down")):
            self.assertEqual(ns.collect(), {})

    def test_an_unscraped_device_is_omitted_not_reported_blank(self):
        """"we never asked" and "we asked and it is down" need different handling, so an
        unknown device must not arrive looking like a measured one."""
        with mock.patch("reports.network_sod.network.device_inventory",
                        return_value=[{"key": "core-switch", "known": False,
                                       "reachable": False, "target": "10.0.0.1"}]):
            self.assertEqual(ns.collect(), {})

    def test_a_reachable_device_prefills_its_row(self):
        with mock.patch("reports.network_sod.network.device_inventory",
                        return_value=[{"key": "core-switch", "known": True,
                                       "reachable": True, "target": "10.0.0.1"}]):
            data = ns.prefilled_checklist()
        row = {c.key: c for c in data["core_wan"]}["ho_core"]
        self.assertTrue(row.live)
        self.assertEqual(row.status, "OK")

    def test_an_unreachable_device_prefills_as_down(self):
        with mock.patch("reports.network_sod.network.device_inventory",
                        return_value=[{"key": "core-switch", "known": True,
                                       "reachable": False, "target": "10.0.0.1"}]):
            data = ns.prefilled_checklist()
        row = {c.key: c for c in data["core_wan"]}["ho_core"]
        self.assertEqual(row.status, "DOWN")


class SodFormRoundTrip(TestCase):
    def test_values_are_taken_verbatim(self):
        """No parsing, no unit inference. This sheet gets signed, and a number the app
        reformatted is a number nobody typed."""
        data = ns.from_post({"loss__dandemutande": " 5.70% ",
                             "result__ho_core": "3 ms",
                             "status__ho_core": "OK"})
        self.assertEqual({c.key: c.loss for c in data["circuits"]}["dandemutande"], "5.70%")
        row = {c.key: c for c in data["core_wan"]}["ho_core"]
        self.assertEqual(row.result, "3 ms")
        self.assertEqual(row.status, "OK")

    def test_an_unknown_status_is_discarded_rather_than_trusted(self):
        data = ns.from_post({"status__ho_core": "TOTALLY FINE"})
        self.assertEqual({c.key: c for c in data["core_wan"]}["ho_core"].status, "")


class SodReportPipeline(TestCase):
    """Tile -> screen -> download, the same shape the other two reports use."""

    def setUp(self):
        self.user = get_user_model().objects.create_user("sod", password="pw12345!")
        self.user.groups.add(Group.objects.get_or_create(name="Network Admin")[0])
        self.client.force_login(self.user)

    def test_the_tile_appears_for_network_admin(self):
        resp = self.client.get(reverse("reports"))
        self.assertContains(resp, "Network Infrastructure SOD Report")

    def test_the_screen_renders_every_section(self):
        with mock.patch("reports.network_sod.network.device_inventory", return_value=[]):
            resp = self.client.get(reverse("network_sod"))
        self.assertEqual(resp.status_code, 200)
        for probe in ("Head-Office Core Switch", "HQ Sophos Firewall", "Cisco Floor Switches",
                      "Dandemutande", "Bulawayo", "Dark Fibre Africa Link", "beam.rbz.co.zw"):
            self.assertContains(resp, probe)

    def test_a_non_network_admin_cannot_reach_it(self):
        other = get_user_model().objects.create_user("nope", password="pw12345!")
        self.client.force_login(other)
        with mock.patch("reports.network_sod.network.device_inventory", return_value=[]):
            resp = self.client.get(reverse("network_sod"))
        self.assertEqual(resp.status_code, 302)

    def test_generate_returns_a_workbook_and_records_the_submission(self):
        from .models import ReportSubmission
        post = {"author": "A. Engineer", "theme": "dark",
                "result__ho_core": "3 ms", "status__ho_core": "OK",
                "loss__dandemutande": "5.70%", "qual__dandemutande": "Average/Poor/Average",
                "rogue__harare": "592"}
        resp = self.client.post(reverse("network_sod_generate"), post)
        self.assertEqual(resp.status_code, 200)
        self.assertIn("spreadsheetml", resp["Content-Type"])
        self.assertIn("Network_Infrastructure_SOD_Report", resp["Content-Disposition"])
        self.assertGreater(len(resp.content), 5000)

        sub = ReportSubmission.objects.latest("id")
        self.assertEqual(sub.author, "A. Engineer")
        self.assertEqual(sub.report_content["kind"], "network_sod")
        # one degraded circuit + one controller with rogue APs
        self.assertEqual(sub.watch_count, 2)
        self.assertEqual(sub.immediate_count, 0)

    def test_a_wholly_blank_submission_still_produces_a_report(self):
        """Blank is a reportable answer — that is the whole point of the sheet."""
        resp = self.client.post(reverse("network_sod_generate"), {"author": "A N Other"})
        self.assertEqual(resp.status_code, 200)
        self.assertGreater(len(resp.content), 5000)

    def test_both_themes_build(self):
        for theme in ("dark", "light"):
            resp = self.client.post(reverse("network_sod_generate"),
                                    {"author": "A", "theme": theme})
            self.assertEqual(resp.status_code, 200, theme)

    def test_the_filename_carries_the_date(self):
        import datetime
        name = ns.sod_report_filename("dark", datetime.datetime(2026, 8, 18))
        self.assertEqual(name, "Network_Infrastructure_SOD_Report_18_Aug_2026.xlsx")
